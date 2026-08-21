#!/usr/bin/env python3
"""Build a compact browser dataset from local Federal IDR PUF files."""

from __future__ import annotations

import csv
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook


APP_ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = APP_ROOT.parent
OUTPUT_PATH = APP_ROOT / "public" / "idr-data.json"
PROCEDURES_PATH = DATA_ROOT / "Procedures_Billing_Codes" / "procedures.json"
CODE_STATS_PATH = DATA_ROOT / "Procedures_Billing_Codes" / "code_stats.json"

ENTITY_NOT_REPORTED = "Not reported in selected source"
REGION_NOT_REPORTED = "N/R"
VALUE_SUPPRESSED = "^"

REQUIRED_COLUMNS = {
    "Service Code",
    "Place of Service Code",
    "Geographical Region",
    "QPA",
    "Provider/Facility Offer",
    "Health Plan/Issuer Offer",
    "Offer Selected from Provider or Issuer",
    "Prevailing Offer",
    "Default Decision",
}


def normalize_code(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    if re.fullmatch(r"\d+\.0", text):
        text = str(int(float(text)))
    if text.isdigit() and len(text) < 5:
        text = text.zfill(5)
    return text.upper()


def clean_text(value: Any, fallback: str = "N/R") -> str:
    if value is None:
        return fallback
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return fallback
    return text


def parse_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if math.isnan(value) if isinstance(value, float) else False:
            return None
        return float(value)
    text = str(value).strip()
    if not text or text in {VALUE_SUPPRESSED, "N/R", "NR", "nan"}:
        return None
    text = text.replace("$", "").replace(",", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def parse_year_of_service(value: Any) -> int | None:
    number = parse_number(value)
    if number is None:
        return None
    year = int(number)
    return year if 1900 <= year <= 2100 else None


def parse_outcome(value: Any) -> int:
    text = clean_text(value, "Other").lower()
    if "split" in text:
        return 2
    if "provider" in text or "facility" in text or "aa provider" in text:
        return 0
    if "plan" in text or "issuer" in text:
        return 1
    return 3


def parse_default(value: Any) -> int:
    text = clean_text(value, "N/R").lower()
    if text == "yes":
        return 1
    if text == "no":
        return 0
    return 2


def parse_source_period(path: Path) -> tuple[int, int]:
    text = str(path).lower()
    year_match = re.search(r"20\d{2}", text)
    quarter_match = re.search(r"q([1-4])", text)
    if not year_match or not quarter_match:
        raise ValueError(f"Could not parse year/quarter from {path}")
    return int(year_match.group(0)), int(quarter_match.group(1))


def is_spine_related_code(code: str) -> bool:
    if not code.isdigit():
        return False
    value = int(code)
    return (
        22000 <= value <= 22899
        or 63000 <= value <= 63799
        or 20930 <= value <= 20939
        or value in {61781, 61782, 61783, 69990}
        or 95860 <= value <= 95941
    )


def discover_sources() -> list[Path]:
    sources: list[Path] = []
    for year in ("2023", "2024"):
        for path in sorted((DATA_ROOT / year).glob("*.xlsx")):
            name = path.name.lower()
            if "puf" in name and "supplemental" not in name:
                sources.append(path)
    for path in sorted((DATA_ROOT / "2025").rglob("*.csv")):
        name = path.name.lower()
        if "qpa" in name and "offer" in name:
            sources.append(path)
    return sorted(sources, key=lambda p: (*parse_source_period(p), str(p)))


def load_code_metadata() -> tuple[list[dict[str, Any]], dict[str, int]]:
    with PROCEDURES_PATH.open(encoding="utf-8") as handle:
        procedures = json.load(handle)

    codes: list[dict[str, Any]] = []
    seen: set[str] = set()
    for procedure in procedures:
        for code_info in procedure.get("codes", []):
            code = normalize_code(code_info.get("code"))
            if not code or code in seen:
                continue
            seen.add(code)
            codes.append(
                {
                    "code": code,
                    "description": clean_text(code_info.get("description"), ""),
                    "status": clean_text(code_info.get("status"), ""),
                    "procedureId": procedure.get("id", ""),
                    "procedureTitle": procedure.get("title", ""),
                    "procedureDescription": procedure.get("description", ""),
                }
            )

    extra_codes: list[dict[str, str]] = []
    if CODE_STATS_PATH.exists():
        with CODE_STATS_PATH.open(encoding="utf-8") as handle:
            stats = json.load(handle).get("codes", {})
        for raw_code in stats:
            code = normalize_code(raw_code)
            if not code or code in seen or not is_spine_related_code(code):
                continue
            seen.add(code)
            code_info = {
                "code": code,
                "description": "Additional spine-related CPT code present in the local IDR billing statistics.",
                "status": "From code_stats.json",
            }
            extra_codes.append(code_info)
            codes.append(
                {
                    **code_info,
                    "procedureId": "SPINE_EXTRA",
                    "procedureTitle": "Spine-related add-on and adjacent codes",
                    "procedureDescription": "Spinal surgery, spinal cord, neuromonitoring, graft, navigation, and microscope codes present in the local billing statistics but not the NHSN procedure mapping.",
                }
            )

    if extra_codes:
        procedures.append(
            {
                "id": "SPINE_EXTRA",
                "title": "Spine-related add-on and adjacent codes",
                "description": "Spinal surgery, spinal cord, neuromonitoring, graft, navigation, and microscope codes present in the local billing statistics but not the NHSN procedure mapping.",
                "codes": sorted(extra_codes, key=lambda item: item["code"]),
            }
        )

    codes.sort(key=lambda item: item["code"])
    return procedures, {item["code"]: index for index, item in enumerate(codes)}


class Dictionary:
    def __init__(self, initial: list[str] | None = None) -> None:
        self.values: list[str] = []
        self.index: dict[str, int] = {}
        for value in initial or []:
            self.get(value)

    def get(self, value: str) -> int:
        if value not in self.index:
            self.index[value] = len(self.values)
            self.values.append(value)
        return self.index[value]


def row_value(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row:
            return row[name]
    return None


def make_record(
    row: dict[str, Any],
    year: int,
    quarter: int,
    code_index: dict[str, int],
    regions: Dictionary,
    entities: Dictionary,
    places: Dictionary,
    line_types: Dictionary,
    initiating_parties: Dictionary,
    modifiers: Dictionary,
) -> list[Any] | None:
    code = normalize_code(row_value(row, "Service Code"))
    code_ix = code_index.get(code)
    if code_ix is None:
        return None

    region = clean_text(row_value(row, "Geographical Region"), REGION_NOT_REPORTED)
    entity = clean_text(row_value(row, "Certified IDR Entity"), ENTITY_NOT_REPORTED)
    place = clean_text(row_value(row, "Place of Service Code"), "N/R")
    line_type = clean_text(row_value(row, "Dispute Line Item Type"), "N/R")
    initiating_party = clean_text(row_value(row, "Initiating Party"), "N/R")
    modifier = clean_text(row_value(row, "Service Code Modifier(s)"), "N/R")

    return [
        code_ix,
        regions.get(region),
        entities.get(entity),
        year,
        quarter,
        places.get(place),
        parse_number(row_value(row, "QPA")),
        parse_number(row_value(row, "Provider/Facility Offer")),
        parse_number(row_value(row, "Health Plan/Issuer Offer")),
        parse_number(row_value(row, "Prevailing Offer")),
        parse_outcome(row_value(row, "Offer Selected from Provider or Issuer")),
        parse_default(row_value(row, "Default Decision")),
        line_types.get(line_type),
        initiating_parties.get(initiating_party),
        modifiers.get(modifier),
        parse_year_of_service(row_value(row, "Year of Service")),
        parse_number(row_value(row, "Cost-Sharing Amount")),
        parse_number(row_value(row, "Initial Payment Amount")),
    ]


def read_csv_source(
    path: Path,
    year: int,
    quarter: int,
    build: Callable[[dict[str, Any], int, int], list[Any] | None],
) -> tuple[list[list[Any]], int]:
    records: list[list[Any]] = []
    scanned = 0
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        missing = REQUIRED_COLUMNS - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"{path} missing columns: {sorted(missing)}")
        for row in reader:
            scanned += 1
            record = build(row, year, quarter)
            if record is not None:
                records.append(record)
            if scanned % 250000 == 0:
                print(f"  scanned {scanned:,} rows; matched {len(records):,}")
    return records, scanned


def read_xlsx_source(
    path: Path,
    year: int,
    quarter: int,
    build: Callable[[dict[str, Any], int, int], list[Any] | None],
) -> tuple[list[list[Any]], int]:
    records: list[list[Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = next(name for name in workbook.sheetnames if name.strip().lower() == "qpa and offers")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [clean_text(value, "") for value in next(rows)]
        missing = REQUIRED_COLUMNS - set(headers)
        if missing:
            raise ValueError(f"{path} missing columns: {sorted(missing)}")
        scanned = 0
        for values in rows:
            scanned += 1
            row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
            record = build(row, year, quarter)
            if record is not None:
                records.append(record)
            if scanned % 250000 == 0:
                print(f"  scanned {scanned:,} rows; matched {len(records):,}")
        return records, scanned
    finally:
        workbook.close()


def main() -> None:
    procedures, code_index = load_code_metadata()
    codes = [None] * len(code_index)
    for code, index in code_index.items():
        codes[index] = code

    regions = Dictionary([REGION_NOT_REPORTED])
    entities = Dictionary([ENTITY_NOT_REPORTED])
    places = Dictionary(["N/R"])
    line_types = Dictionary(["N/R"])
    initiating_parties = Dictionary(["N/R"])
    modifiers = Dictionary(["N/R"])

    sources = discover_sources()
    records: list[list[Any]] = []
    source_summaries: list[dict[str, Any]] = []

    def build(row: dict[str, Any], year: int, quarter: int) -> list[Any] | None:
        return make_record(
            row,
            year,
            quarter,
            code_index,
            regions,
            entities,
            places,
            line_types,
            initiating_parties,
            modifiers,
        )

    for path in sources:
        year, quarter = parse_source_period(path)
        print(f"Reading {year} Q{quarter}: {path.relative_to(DATA_ROOT)}")
        if path.suffix.lower() == ".csv":
            new_records, scanned = read_csv_source(path, year, quarter, build)
        else:
            new_records, scanned = read_xlsx_source(path, year, quarter, build)
        records.extend(new_records)
        source_summaries.append(
            {
                "year": year,
                "quarter": quarter,
                "file": str(path.relative_to(DATA_ROOT)),
                "scannedRows": scanned,
                "matchedRows": len(new_records),
            }
        )
        print(f"  done; matched {len(new_records):,} of {scanned:,}")

    code_objects = [None] * len(code_index)
    for procedure in procedures:
        for code_info in procedure.get("codes", []):
            code = normalize_code(code_info.get("code"))
            index = code_index.get(code)
            if index is None:
                continue
            code_objects[index] = {
                "code": code,
                "description": clean_text(code_info.get("description"), ""),
                "status": clean_text(code_info.get("status"), ""),
                "procedureId": procedure.get("id", ""),
                "procedureTitle": procedure.get("title", ""),
                "procedureDescription": procedure.get("description", ""),
            }

    payload = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "recordColumns": [
            "code",
            "region",
            "entity",
            "year",
            "quarter",
            "place",
            "qpa",
            "providerOffer",
            "issuerOffer",
            "prevailingOffer",
            "outcome",
            "defaultDecision",
            "lineType",
            "initiatingParty",
            "modifier",
            "yearOfService",
            "costSharingAmount",
            "initialPaymentAmount",
        ],
        "outcomes": ["provider", "issuer", "split", "other"],
        "defaultValues": ["no", "yes", "unknown"],
        "dictionaries": {
            "codes": code_objects,
            "regions": regions.values,
            "entities": entities.values,
            "places": places.values,
            "lineTypes": line_types.values,
            "initiatingParties": initiating_parties.values,
            "modifiers": modifiers.values,
        },
        "procedures": procedures,
        "sources": source_summaries,
        "records": records,
        "notes": [
            "Dollar fields are null when CMS suppressed a cell, usually shown as '^' in the source files.",
            "Certified IDR Entity is available in the local PUF data beginning with 2025 Q3; older rows are marked as not reported.",
            "Recommendation modeling excludes default decisions when enough contested comparable rows are available.",
        ],
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_PATH.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, separators=(",", ":"))
    print(f"Wrote {OUTPUT_PATH.relative_to(APP_ROOT)}")
    print(f"Records: {len(records):,}; size: {OUTPUT_PATH.stat().st_size / 1024 / 1024:.1f} MB")


if __name__ == "__main__":
    main()
